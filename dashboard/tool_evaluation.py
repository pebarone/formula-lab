"""Paired formula evaluation; Excel execution is diagnostic, not the answer oracle."""

import base64
import hashlib
import io
import json
import math
import re
import time

from benchmark import run_baseline as baseline
from benchmark.scoring import extract_prediction, formulas_match
from dashboard.config import ROOT, read_json, write_json
from dashboard.runtime import Cancelled


def validate_tool_eval(raw):
    if not isinstance(raw, dict) or set(raw) - {"test_count", "max_tool_calls"}:
        raise ValueError("Parâmetros do teste de ferramentas inválidos.")
    result = {"test_count": 10, "max_tool_calls": 4, **raw}
    for name, maximum in (("test_count", 500), ("max_tool_calls", 4)):
        value = result[name]
        if isinstance(value, bool) or not isinstance(value, int) or not 1 <= value <= maximum:
            raise ValueError(f"{name} deve ser um inteiro de 1 a {maximum}.")
    return result


def _sample_value(text):
    # Dates, percentages and ambiguous formats remain the literal dataset text.
    if re.fullmatch(r"[+-]?(?:0|[1-9]\d*)(?:\.\d+)?", text) and len(text.replace(".", "").lstrip("+-0")) <= 15:
        value = float(text) if "." in text else int(text)
        if math.isfinite(value):
            return value
    return text


def sample_workbook(context):
    """Materialize exactly the visible sample, without calculating any formulas."""
    from openpyxl import Workbook
    from openpyxl.utils import column_index_from_string

    lines = context.splitlines()
    if not lines or not lines[0].startswith("Sheet: ") or "Sample:" not in lines:
        raise ValueError("O exemplo não contém aba e amostra de células reconhecíveis.")
    name = lines[0][7:]
    if not name or len(name) > 31 or any(char in name for char in ":\\/?*[]"):
        raise ValueError("Nome de aba inválido no exemplo.")
    split = lines.index("Sample:")
    columns = []
    book = Workbook()
    sheet = book.active
    sheet.title = name
    for line in lines[1:split]:
        match = re.fullmatch(r"([A-Z]{1,2}): (.*)", line)
        if not match:
            raise ValueError("Cabeçalho da amostra inválido.")
        column = column_index_from_string(match[1])
        if column > 40 or column in columns:
            raise ValueError("Colunas da amostra inválidas ou duplicadas.")
        columns.append(column)
        cell = sheet.cell(1, column, match[2])
        cell.data_type = "s"
    if not columns or len(lines[split + 1:]) > 500:
        raise ValueError("Amostra vazia ou grande demais.")
    seen = set()
    for line in lines[split + 1:]:
        parts = line.split(" | ")
        if len(parts) != len(columns) + 1 or not parts[0].isdigit():
            raise ValueError("Linha da amostra ambígua ou inválida.")
        row = int(parts[0])
        if not 2 <= row <= 501 or row in seen:
            raise ValueError("Endereço de linha inválido ou duplicado na amostra.")
        seen.add(row)
        for column, text in zip(columns, parts[1:]):
            value = _sample_value(text)
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # A leading '=' is data, never an imported formula.
    output = io.BytesIO()
    book.save(output)
    return name, output.getvalue()


def compare_example(row, generate, workbook_id, check_formula, max_calls=4, check_cancel=None):
    """Keep the expected answer out of generation and all tool feedback."""
    from dashboard.excel_tools import run_tool_conversation

    messages = baseline.row_messages(row)[1:]
    results = {}
    for condition in ("without_tools", "with_tools"):
        started = time.perf_counter()
        turns = []

        def turn(history, tools=None):
            generated = generate(history, tools)
            turns.append({key: generated[key] for key in (
                "generated_tokens", "generation_seconds", "generation_mode", "prompt_tokens",
            ) if key in generated})
            return generated["text"]

        try:
            if condition == "with_tools":
                answer = run_tool_conversation(messages, turn, workbook_id, max_calls=max_calls, check_cancel=check_cancel)
            else:
                answer = {"text": turn(messages), "tool_trace": [], "tool_calls": 0}
            predicted = extract_prediction(answer["text"])
            entry = {**answer, "predicted": predicted, "match": formulas_match(row["expected"], predicted)}
            if predicted.startswith("="):
                try:
                    execution = check_formula(predicted)
                    entry.update(execution=execution, execution_ok=bool(execution.get("ok")))
                except Cancelled:
                    raise
                except Exception as exc:
                    entry.update(execution_ok=False, execution_error=str(exc))
            else:
                entry.update(execution_ok=False, execution_error="A resposta não contém fórmula.")
        except Cancelled:
            raise
        except Exception as exc:
            entry = {"predicted": "", "match": False, "execution_ok": False, "error": str(exc), "tool_calls": 0}
        entry.update(turns=turns, elapsed_seconds=time.perf_counter() - started)
        results[condition] = entry
    return {"id": row["id"], "instruction": row["instruction"], "expected": row["expected"], **results}


def summarize(items, requested_total, metadata):
    total = len(items)
    result = {**metadata, "total": total, "requested_total": requested_total,
              "metric": "normalized_exact_match", "items": items}
    for condition in ("without_tools", "with_tools"):
        correct = sum(bool(item[condition]["match"]) for item in items)
        result[condition] = {
            "correct": correct, "accuracy": correct / total if total else 0,
            "execution_ok": sum(bool(item[condition]["execution_ok"]) for item in items),
            "tool_calls": sum(item[condition].get("tool_calls", 0) for item in items),
            "examples_using_tools": sum(bool(item[condition].get("tool_calls")) for item in items),
        }
    result["comparison"] = {
        "delta_pp": (result["with_tools"]["accuracy"] - result["without_tools"]["accuracy"]) * 100,
        "gained": sum(not item["without_tools"]["match"] and item["with_tools"]["match"] for item in items),
        "lost": sum(item["without_tools"]["match"] and not item["with_tools"]["match"] for item in items),
    }
    return result


def evaluate_tools(request, tokenizer, directory, events):
    from dashboard.attachments import save_attachment
    from dashboard.chat_backend import generate_chat_turn
    from dashboard.chat_engine import Resident
    from dashboard.training import limit_training_memory, training_memory
    from dashboard.workbooks import create_workbook, test_formula

    settings = validate_tool_eval(request.get("tool_eval", {}))
    data = ROOT / "dataset/canonical/test.jsonl"
    all_rows = baseline.load_jsonl(data, None)
    rows = all_rows[:settings["test_count"]]
    if len(rows) != settings["test_count"]:
        raise ValueError("O dataset não contém a quantidade de testes pedida.")
    memory = limit_training_memory(request["config"], events, operation="teste de ferramentas")
    resident = Resident()
    model = resident.model_for(request.get("adapter_run"), events, request.get("decensor_run"))
    metadata = {
        "adapter_run": request.get("adapter_run"), "decensor_run": request.get("decensor_run"),
        "data_sha256": hashlib.sha256(data.read_bytes()).hexdigest(),
        "evaluation_identity": baseline.evaluation_identity(rows),
        "settings": settings, "engine": "Microsoft Excel", "status": "running",
        "harness_sha256": {name: hashlib.sha256((ROOT / "dashboard" / name).read_bytes()).hexdigest()
                           for name in ("excel_tools.py", "tool_evaluation.py", "chat_backend.py", "excel_engine.py", "excel_worker.py")},
        "scope": "Only the visible context sample is materialized; Excel success is not semantic correctness.",
        "generation": {"do_sample": False, "without_tools_max_new_tokens": 128,
                       "with_tools_max_new_tokens_per_turn": 512, "max_tool_rounds": 4},
        "memory": memory,
    }
    items = []
    for index, row in enumerate(rows):
        events.check()
        events.emit(phase="Comparando ferramentas", step=index, total_steps=len(rows), phase_unit="exemplos")
        try:
            sheet, content = sample_workbook(row["context_text"])
            attachment = save_attachment("amostra.xlsx", base64.b64encode(content).decode("ascii"))
            pristine = create_workbook(attachment["id"])
            working = create_workbook(attachment["id"])

            def generate(history, tools):
                events.check()
                return generate_chat_turn(history, tokenizer, model, resident.decoder, events,
                                          tools=tools, max_new_tokens=128 if tools is None else 512,
                                          adapter_run=request.get("adapter_run"), decensor_run=request.get("decensor_run"))

            def check_formula(formula):
                events.check()
                # Always score execution on the original inputs, even if the model edited its working copy.
                return test_formula(pristine["id"], sheet, "AZ1", formula)

            item = compare_example(row, generate, working["id"], check_formula, settings["max_tool_calls"], events.check)
            item.update(workbook_id=working["id"], input_workbook_id=pristine["id"])
        except Cancelled:
            raise
        except Exception as exc:
            failure = {"predicted": "", "match": False, "execution_ok": False, "tool_calls": 0, "error": str(exc)}
            item = {"id": row["id"], "instruction": row["instruction"], "expected": row["expected"],
                    "without_tools": dict(failure), "with_tools": dict(failure), "setup_error": str(exc)}
        items.append(item)
        with (directory / "tool_predictions.jsonl").open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(item, ensure_ascii=False, allow_nan=False) + "\n")
        report = summarize(items, len(rows), metadata)
        summary = {key: value for key, value in report.items() if key != "items"}
        write_json(directory / "tool_summary.json", summary)
        write_json(directory / "tool_report.json", {**summary, "items": []})
        events.emit(phase="Comparando ferramentas", step=index + 1, total_steps=len(rows),
                    phase_unit="exemplos", **training_memory())
    report.update(status="completed", memory={**memory, **training_memory()},
                  environment=read_json(directory / "environment.json", {}))
    write_json(directory / "tool_report.json", report)
    write_json(directory / "tool_summary.json", {key: value for key, value in report.items() if key != "items"})
    events.emit(phase="Comparação de ferramentas concluída", step=len(rows), total_steps=len(rows))
