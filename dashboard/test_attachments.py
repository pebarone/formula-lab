import base64
import io
import json
import re
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from uuid import uuid4

from openpyxl import Workbook

from dashboard.attachments import _check_xlsx_archive, load_attachments, save_attachment
from dashboard.config import read_json
from dashboard.server import Handler, Jobs


def encoded(value):
    return base64.b64encode(value).decode("ascii")


class AttachmentTests(unittest.TestCase):
    def test_csv_unicode_formulas_and_original_are_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            content = "Produto;Preço\nCafé;=SUM(1,2)\n".encode("utf-8-sig")
            result = save_attachment("vendas.csv", encoded(content), runs_root=temp)

            self.assertIn('A2="Café"', result["context"])
            self.assertIn('B2="=SUM(1,2)"', result["context"])
            self.assertIn("Faixa amostrada: A1:B2", result["context"])
            self.assertEqual(result["statistics"]["formulas_sampled"], 1)
            self.assertFalse(result["statistics"]["truncated"])
            self.assertEqual(load_attachments([result["id"]], runs_root=temp)[0]["context"], result["context"])
            self.assertEqual((Path(temp) / "attachments" / result["id"] / "original.csv").read_bytes(), content)

    def test_csv_row_and_column_limits_are_reported(self):
        with tempfile.TemporaryDirectory() as temp, patch("dashboard.attachments.MAX_CSV_ROWS", 2), \
                patch("dashboard.attachments.MAX_COLUMNS", 1):
            content = b"a,b\n1,2\n3,4\n"
            result = save_attachment("wide.csv", encoded(content), runs_root=temp)

        self.assertTrue(result["statistics"]["truncated"])
        self.assertTrue(any("linhas" in warning for warning in result["warnings"]))
        self.assertTrue(any("colunas" in warning for warning in result["warnings"]))
        self.assertIn("A1:A2", result["context"])

    def test_invalid_name_encoding_and_decoded_size_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            for name, data in (("../sales.csv", b"a,b\n"), ("sales.xlsm", b"x"), ("bad.csv", b"\xff")):
                with self.subTest(name=name), self.assertRaises(ValueError):
                    save_attachment(name, encoded(data), runs_root=temp)
            with self.assertRaises(ValueError):
                save_attachment("sales.csv", "%%%", runs_root=temp)
            with patch("dashboard.attachments.MAX_FILE_BYTES", 3), self.assertRaises(ValueError):
                save_attachment("sales.csv", encoded(b"four"), runs_root=temp)

    def test_unreadable_and_encrypted_xlsx_files_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "corrompido"):
            save_attachment("bad.xlsx", encoded(b"not a zip file"), runs_root=temp)

        archive = MagicMock()
        archive.__enter__.return_value = archive
        archive.infolist.return_value = [SimpleNamespace(
            filename="xl/workbook.xml", flag_bits=1, file_size=10, compress_size=10,
        )]
        with patch("dashboard.attachments.zipfile.ZipFile", return_value=archive), \
                self.assertRaisesRegex(ValueError, "protegidas por senha"):
            _check_xlsx_archive(b"encrypted fixture")

    def test_xlsx_multiple_sheets_keep_addresses_and_formula_text(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "Dados"
        first["A1"] = "Região"
        first["A2"] = "São Paulo"
        first["B2"] = "=SUM(1,2)"
        second = workbook.create_sheet("Resumo")
        second["A1"] = "Total"
        second["B1"] = "=Dados!B2"
        buffer = io.BytesIO()
        workbook.save(buffer)

        with tempfile.TemporaryDirectory() as temp:
            result = save_attachment("planilha.xlsx", encoded(buffer.getvalue()), runs_root=temp)

        self.assertEqual(result["sheets"], ["Dados", "Resumo"])
        self.assertIn('Aba "Dados": amostra A1:B2', result["context"])
        self.assertIn('A2="São Paulo"', result["context"])
        self.assertIn('B2="=SUM(1,2)"', result["context"])
        self.assertIn('Aba "Resumo"', result["context"])
        self.assertEqual(result["statistics"]["formulas_sampled"], 2)

    def test_xlsx_zip_bombs_and_macros_are_rejected_before_openpyxl(self):
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("xl/workbook.xml", "workbook")
            archive.writestr("xl/large.xml", b"0" * (4 * 1024 * 1024))
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "compactação"):
            save_attachment("bomb.xlsx", encoded(bomb.getvalue()), runs_root=temp)

        macro = io.BytesIO()
        with zipfile.ZipFile(macro, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("xl/workbook.xml", "workbook")
            archive.writestr("xl/vbaProject.bin", "macro")
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "macros"):
            save_attachment("macro.xlsx", encoded(macro.getvalue()), runs_root=temp)

        xml = io.BytesIO()
        with zipfile.ZipFile(xml, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("xl/workbook.xml", '<!DOCTYPE workbook [<!ENTITY x "expanded">]><workbook>&x;</workbook>')
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "Declarações XML"):
            save_attachment("entity.xlsx", encoded(xml.getvalue()), runs_root=temp)

        utf16 = io.BytesIO()
        with zipfile.ZipFile(utf16, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("xl/workbook.xml", '<!DOCTYPE workbook [<!ENTITY x "expanded">]><workbook>&x;</workbook>'.encode("utf-16"))
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "Declarações XML"):
            save_attachment("entity-utf16.xlsx", encoded(utf16.getvalue()), runs_root=temp)

    def test_xlsx_exact_total_row_limit_does_not_report_omitted_data(self):
        for sizes, truncated in (((250, 250), False), ((250, 250, 1), True), ((250, 251), True)):
            with self.subTest(sheet_rows=sizes), tempfile.TemporaryDirectory() as temp:
                workbook = Workbook()
                for index, rows in enumerate(sizes):
                    sheet = workbook.active if index == 0 else workbook.create_sheet()
                    for value in range(rows):
                        sheet.append([value])
                buffer = io.BytesIO()
                workbook.save(buffer)
                result = save_attachment("boundary.xlsx", encoded(buffer.getvalue()), runs_root=temp)
                self.assertEqual(result["statistics"]["rows_sampled"], 500)
                self.assertEqual(result["statistics"]["truncated"], truncated)
                self.assertEqual(bool(result["warnings"]), truncated)

    def test_xlsx_without_optional_dimensions_reads_a_bounded_sample_and_warns(self):
        workbook = Workbook()
        workbook.active["A1"] = "Cabeçalho"
        workbook.active["C4"] = "Dentro da amostra"
        workbook.active["D6"] = 42
        original = io.BytesIO()
        workbook.save(original)
        modified = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w", zipfile.ZIP_DEFLATED) as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb"<dimension\b[^>]*/>", b"", content)
                target.writestr(member, content)
        with tempfile.TemporaryDirectory() as temp:
            result = save_attachment("unknown-dimensions.xlsx", encoded(modified.getvalue()), runs_root=temp)
        self.assertIn('C4="Dentro da amostra"', result["context"])
        self.assertIn("D6=42", result["context"])
        self.assertTrue(result["statistics"]["truncated"])
        self.assertTrue(any("dimensões ausentes" in warning for warning in result["warnings"]))
        self.assertLessEqual(result["statistics"]["rows_sampled"], 250)

    def test_attachment_ids_are_bounded_and_server_resolved(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaises(ValueError):
                load_attachments(["../somewhere"], runs_root=temp)
            with self.assertRaises(ValueError):
                load_attachments([str(uuid4())] * 4, runs_root=temp)
            with self.assertRaises(ValueError):
                load_attachments(["not-a-uuid", "not-a-uuid"], runs_root=temp)

    def test_xlsx_undersized_declared_dimensions_are_rejected(self):
        workbook = Workbook()
        workbook.active["A1"] = "Visível"
        workbook.active["C4"] = "Não pode sumir"
        workbook.active["ZZ1000"] = 42
        original = io.BytesIO()
        workbook.save(original)
        modified = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w", zipfile.ZIP_DEFLATED) as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb"<dimension\b[^>]*/>", b'<dimension ref="A1:A1"/>', content)
                target.writestr(member, content)
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "dimensões inconsistentes"):
            save_attachment("wrong-dimensions.xlsx", encoded(modified.getvalue()), runs_root=temp)

    def test_jobs_start_returns_canonical_messages_and_stored_metadata(self):
        with tempfile.TemporaryDirectory() as temp:
            attachment = save_attachment("data.csv", encoded(b"name,value\nA,1\n"), runs_root=temp)
            with patch("dashboard.server.RUNS", Path(temp)), patch("dashboard.server.dispatch_chat") as dispatch:
                result = Jobs().start({
                    "kind": "chat",
                    "messages": [{"role": "user", "content": "Summarize this table"}],
                    "attachment_ids": [attachment["id"]],
                })

            self.assertIn("Summarize this table", result["messages"][-1]["content"])
            self.assertIn('A1="name"', result["messages"][-1]["content"])
            self.assertEqual(result["attachments"][0]["name"], "data.csv")
            self.assertNotIn("context", result["attachments"][0])
            request = read_json(Path(temp) / result["id"] / "request.json")
            self.assertEqual(request["messages"], result["messages"])
            self.assertEqual(request["attachments"], [{"id": attachment["id"], "name": "data.csv", "kind": "csv"}])
            dispatch.assert_called_once_with(result["id"])

    def test_invalid_attachment_never_dispatches_or_spawns_worker(self):
        with tempfile.TemporaryDirectory() as temp, patch("dashboard.server.RUNS", Path(temp)), \
                patch("dashboard.server.dispatch_chat") as dispatch, patch("dashboard.server.subprocess.Popen") as spawn:
            with self.assertRaises(ValueError):
                Jobs().start({"kind": "chat", "messages": [{"role": "user", "content": "question"}],
                              "attachment_ids": [str(uuid4())]})
            with self.assertRaises(ValueError):
                Jobs().start({"kind": "train", "attachment_ids": []})
            dispatch.assert_not_called()
            spawn.assert_not_called()


class AttachmentEndpointTests(unittest.TestCase):
    def setUp(self):
        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.httpd.token = "test-session-token"
        self.httpd.jobs = None
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.httpd.shutdown()
        self.thread.join(timeout=2)
        self.httpd.server_close()

    def post(self, path, body, token=True):
        url = f"http://127.0.0.1:{self.httpd.server_port}{path}"
        headers = {"Content-Type": "application/json"}
        if token:
            headers["X-Session-Token"] = self.httpd.token
        request = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"), headers=headers, method="POST")
        return urllib.request.urlopen(request, timeout=3)

    def test_upload_route_uses_session_token_and_persists_file(self):
        with tempfile.TemporaryDirectory() as temp, patch("dashboard.server.RUNS", Path(temp)):
            response = self.post("/api/attachments", {"name": "data.tsv", "content_base64": encoded("chave\tvalor\né\t1\n".encode())})
            self.assertEqual(response.status, 201)
            result = json.load(response)
            self.assertEqual(result["kind"], "tsv")
            self.assertTrue((Path(temp) / "attachments" / result["id"] / "original.tsv").is_file())

            with self.assertRaises(urllib.error.HTTPError) as denied:
                self.post("/api/attachments", {"name": "data.tsv", "content_base64": encoded(b"a\tb\n")}, token=False)
            self.assertEqual(denied.exception.code, 403)

    def test_chat_start_has_a_separate_larger_json_limit(self):
        self.httpd.jobs = type("JobsStub", (), {"start": lambda _self, body: {"chars": len(body["messages"][0]["content"])}})()
        body = {"kind": "chat", "messages": [{"role": "user", "content": "x" * 70_000}], "attachment_ids": []}
        response = self.post("/api/start", body)
        self.assertEqual(json.load(response)["chars"], 70_000)


if __name__ == "__main__":
    unittest.main()
