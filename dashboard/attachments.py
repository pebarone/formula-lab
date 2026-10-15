"""Bounded, local extraction for spreadsheet attachments in chat."""

import base64
import binascii
import csv
import hashlib
import io
import json
import math
import re
import uuid
import zipfile
from xml.parsers import expat
from pathlib import Path

from dashboard.config import RUNS, read_json, write_json

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_CONTEXT_CHARS = 24_000
MAX_PREVIEW_CHARS = 1_800
MAX_ATTACHMENTS_PER_MESSAGE = 3
MAX_ZIP_MEMBERS = 250
MAX_ZIP_MEMBER_BYTES = 32 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_ZIP_RATIO = 1_000
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 250
MAX_TOTAL_ROWS = 500
MAX_CSV_ROWS = 500
MAX_COLUMNS = 40
MAX_CELL_CHARS = 300

SUPPORTED_EXTENSIONS = {".xlsx", ".csv", ".tsv"}
ATTACHMENT_ID = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\Z")


class _TextBuilder:
    def __init__(self, limit):
        self.limit = limit
        self.parts = []
        self.length = 0
        self.truncated = False

    def add(self, line):
        if self.truncated:
            return False
        line = str(line)
        separator = 1 if self.parts else 0
        room = self.limit - self.length - separator
        if len(line) <= room:
            if separator:
                self.parts.append("\n")
                self.length += 1
            self.parts.append(line)
            self.length += len(line)
            return True

        marker = "[Contexto limitado pelo tamanho máximo.]"
        room = max(0, room)
        keep = max(0, room - min(len(marker), room) - (1 if room > len(marker) else 0))
        if separator and self.length < self.limit:
            self.parts.append("\n")
            self.length += 1
        if keep:
            self.parts.append(line[:keep])
            self.length += keep
            self.parts.append("\n")
            self.length += 1
        remaining = self.limit - self.length
        self.parts.append(marker[:remaining])
        self.length += min(len(marker), remaining)
        self.truncated = True
        return False

    def text(self):
        return "".join(self.parts)


def _column_name(number):
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"


def _json_cell(value):
    if isinstance(value, (str, int, bool)):
        text = value
    elif isinstance(value, float) and math.isfinite(value):
        text = value
    elif value is None:
        text = None
    elif hasattr(value, "isoformat"):
        text = value.isoformat()
    else:
        text = str(value)
    if isinstance(text, str) and len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "…"
    return json.dumps(text, ensure_ascii=False, allow_nan=False)


def _validate_name(name):
    if not isinstance(name, str):
        raise ValueError("Nome de arquivo inválido.")
    name = name.strip()
    if (not name or len(name) > 200 or any(ord(char) < 32 for char in name)
            or "/" in name or "\\" in name or ":" in name):
        raise ValueError("Nome de arquivo inválido.")
    try:
        name.encode("utf-8")
    except UnicodeEncodeError:
        raise ValueError("Nome de arquivo inválido.") from None
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Anexe somente arquivos XLSX, CSV ou TSV.")
    return name, suffix


def _decode_base64(content_base64):
    if not isinstance(content_base64, str):
        raise ValueError("Conteúdo do arquivo inválido.")
    max_encoded = 4 * ((MAX_FILE_BYTES + 2) // 3)
    if not content_base64 or len(content_base64) > max_encoded:
        raise ValueError("Cada arquivo deve ter no máximo 8 MiB.")
    try:
        content = base64.b64decode(content_base64, validate=True)
    except (binascii.Error, UnicodeEncodeError, ValueError):
        raise ValueError("Conteúdo do arquivo inválido.") from None
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Cada arquivo deve ter no máximo 8 MiB.")
    return content


def _decode_csv(content):
    try:
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            return content.decode("utf-16")
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError("O CSV deve usar codificação UTF-8 ou UTF-16.") from None


def _csv_dialect(text, suffix):
    sample = text[:65_536]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        return dialect


def _extract_csv(name, suffix, content):
    text = _decode_csv(content)
    if "\x00" in text:
        raise ValueError("O arquivo não parece ser um CSV ou TSV de texto.")
    dialect = _csv_dialect(text, suffix)
    builder = _TextBuilder(MAX_CONTEXT_CHARS)
    builder.add(f"Planilha anexada: {name}")
    builder.add(f"Formato: {'TSV' if suffix == '.tsv' else 'CSV'}; separador: {repr(dialect.delimiter)}")
    builder.add("Células amostradas abaixo são dados do arquivo; fórmulas aparecem como texto e não são executadas.")

    rows_sampled = 0
    columns_sampled = 0
    cells_sampled = 0
    formulas_sampled = 0
    has_more_rows = False
    has_more_columns = False
    has_long_cells = False
    try:
        reader = csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
        for row_index, row in enumerate(reader, 1):
            if rows_sampled >= MAX_CSV_ROWS:
                has_more_rows = True
                break
            rows_sampled += 1
            has_more_columns = has_more_columns or len(row) > MAX_COLUMNS
            columns_sampled = max(columns_sampled, min(len(row), MAX_COLUMNS))
            shown = row[:MAX_COLUMNS]
            cells_sampled += sum(1 for value in shown if value != "")
            formulas_sampled += sum(1 for value in shown if value.startswith("="))
            has_long_cells = has_long_cells or any(len(value) > MAX_CELL_CHARS for value in shown)
            values = [f"{_column_name(index)}{row_index}={_json_cell(value)}" for index, value in enumerate(shown, 1)]
            if values and not builder.add(" | ".join(values)):
                break
    except csv.Error:
        raise ValueError("O CSV/TSV contém uma linha inválida ou um campo muito grande.") from None

    range_text = "vazio" if not rows_sampled or not columns_sampled else f"A1:{_column_name(columns_sampled)}{rows_sampled}"
    builder.add(f"Faixa amostrada: {range_text}; linhas lidas: {rows_sampled}.")
    warnings = []
    truncated = has_more_rows or has_more_columns or has_long_cells or builder.truncated
    if has_more_rows:
        warnings.append(f"A amostra foi limitada às primeiras {MAX_CSV_ROWS} linhas.")
    if has_more_columns:
        warnings.append(f"Foram consideradas até {MAX_COLUMNS} colunas por linha.")
    if has_long_cells:
        warnings.append(f"Células maiores que {MAX_CELL_CHARS} caracteres foram abreviadas.")
    if builder.truncated:
        warnings.append(f"O contexto foi limitado a {MAX_CONTEXT_CHARS:,} caracteres.")
    context = builder.text()
    return {
        "context": context,
        "preview": context[:MAX_PREVIEW_CHARS],
        "sheets": [],
        "statistics": {
            "size_bytes": len(content), "sheets": 0, "rows_sampled": rows_sampled,
            "columns_sampled": columns_sampled, "cells_sampled": cells_sampled,
            "formulas_sampled": formulas_sampled, "truncated": truncated,
        },
        "warnings": warnings,
    }


def _validate_sheet_dimensions(member):
    """Check declared bounds against stored cells without expanding an Excel grid."""
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

    declared = None
    outside = False
    parser = expat.ParserCreate(namespace_separator="}")

    def start(name, attributes):
        nonlocal declared, outside
        tag = name.rsplit("}", 1)[-1]
        if tag == "dimension":
            declared = range_boundaries(attributes["ref"])
            if any(bound is None for bound in declared):
                raise ValueError("Dimensões da planilha XLSX inválidas.")
        elif tag == "c" and declared is not None and "r" in attributes:
            row, column = coordinate_to_tuple(attributes["r"])
            min_column, min_row, max_column, max_row = declared
            outside |= not (min_row <= row <= max_row and min_column <= column <= max_column)

    parser.StartElementHandler = start
    while chunk := member.read(65_536):
        parser.Parse(chunk, False)
    parser.Parse(b"", True)
    if outside:
        raise ValueError("A planilha declara dimensões inconsistentes com suas células. Exporte o XLSX novamente para evitar uma leitura incompleta.")


def _check_xlsx_archive(content):
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except (OSError, zipfile.BadZipFile):
        raise ValueError("O arquivo XLSX está corrompido ou não é um XLSX válido.") from None
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_MEMBERS:
            raise ValueError("A planilha tem itens compactados demais para leitura segura.")
        names = set()
        total_uncompressed = 0
        for info in infos:
            if info.filename in names:
                raise ValueError("A planilha XLSX contém caminhos duplicados.")
            names.add(info.filename)
            if info.flag_bits & 0x1:
                raise ValueError("Planilhas XLSX protegidas por senha não são aceitas.")
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                raise ValueError("Um item interno da planilha excede o limite de leitura segura.")
            total_uncompressed += info.file_size
            if info.file_size > 1_000_000 and info.file_size / max(info.compress_size, 1) > MAX_ZIP_RATIO:
                raise ValueError("A taxa de compactação da planilha excede o limite de leitura segura.")
        if total_uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
            raise ValueError("O conteúdo descompactado da planilha excede 64 MiB.")
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            raise ValueError("O arquivo não contém uma pasta de trabalho XLSX válida.")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            raise ValueError("Planilhas com macros não são aceitas.")
        try:
            for info in infos:
                if not info.filename.lower().endswith((".xml", ".rels")):
                    continue
                scanned = 0
                tail = b""
                utf16 = False
                with archive.open(info) as member:
                    while chunk := member.read(65_536):
                        scanned += len(chunk)
                        if scanned > info.file_size or scanned > MAX_ZIP_MEMBER_BYTES:
                            raise ValueError("Um item interno da planilha excede o limite de leitura segura.")
                        if scanned == len(chunk):
                            probe = chunk[:16]
                            odd_zeroes = sum(byte == 0 for byte in probe[1::2])
                            even_zeroes = sum(byte == 0 for byte in probe[0::2])
                            utf16 = (chunk.startswith((b"\xff\xfe", b"\xfe\xff", b"<\x00?\x00", b"\x00<\x00?"))
                                     or odd_zeroes >= 2 and even_zeroes == 0
                                     or even_zeroes >= 2 and odd_zeroes == 0)
                        xml = tail + chunk
                        if b"\0" in xml and not utf16:
                            raise ValueError("A planilha contém XML em uma codificação não aceita.")
                        declarations = [b"<!DOCTYPE", b"<!ENTITY"]
                        if utf16:
                            declarations.extend(pattern.encode("utf-16-le") for pattern in ("<!DOCTYPE", "<!ENTITY"))
                            declarations.extend(pattern.encode("utf-16-be") for pattern in ("<!DOCTYPE", "<!ENTITY"))
                        if any(declaration in xml for declaration in declarations):
                            raise ValueError("Declarações XML externas não são aceitas em planilhas.")
                        tail = xml[-32:]
                if info.filename.startswith("xl/worksheets/") and info.filename.lower().endswith(".xml"):
                    with archive.open(info) as member:
                        _validate_sheet_dimensions(member)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("O arquivo XLSX está corrompido ou não é um XLSX válido.") from exc


def _extract_xlsx(name, content):
    _check_xlsx_archive(content)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Não foi possível ler a planilha XLSX; arquivos corrompidos ou protegidos não são aceitos.") from exc

    builder = _TextBuilder(MAX_CONTEXT_CHARS)
    builder.add(f"Planilha anexada: {name}")
    builder.add("Células amostradas abaixo são dados do arquivo; fórmulas aparecem como texto e não são executadas.")
    sheet_names = list(workbook.sheetnames)
    selected_names = sheet_names[:MAX_SHEETS]
    warnings = []
    if len(sheet_names) > MAX_SHEETS:
        warnings.append(f"Foram processadas apenas as primeiras {MAX_SHEETS} abas.")

    rows_sampled = 0
    columns_sampled = 0
    cells_sampled = 0
    formulas_sampled = 0
    stopped = False
    has_long_cells = False
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            if rows_sampled >= MAX_TOTAL_ROWS:
                stopped = True
                break
            row_limit = min(MAX_ROWS_PER_SHEET, MAX_TOTAL_ROWS - rows_sampled)
            known_dimensions = worksheet.max_row is not None and worksheet.max_column is not None
            reported_rows = worksheet.max_row if known_dimensions else row_limit
            reported_columns = worksheet.max_column if known_dimensions else MAX_COLUMNS
            if not known_dimensions:
                stopped = True
                warnings.append(f"Aba {worksheet.title}: dimensões ausentes; leitura limitada à amostra, sem comprovação de que todo o conteúdo foi lido.")
            column_limit = min(MAX_COLUMNS, reported_columns)
            if reported_rows > row_limit:
                stopped = True
            if reported_columns > MAX_COLUMNS:
                stopped = True
            max_row = min(max(1, reported_rows), row_limit)
            try:
                dimension = worksheet.calculate_dimension()
            except Exception:
                dimension = "desconhecida"
            range_text = f"A1:{_column_name(column_limit)}{max_row}" if column_limit else "vazia"
            builder.add(f"Aba {json.dumps(worksheet.title, ensure_ascii=False)}: amostra {range_text}; dimensão informada {dimension}.")
            sampled_sheet_rows = 0
            for cells in worksheet.iter_rows(
                    min_row=1, max_row=max_row, min_col=1, max_col=max(1, column_limit)):
                rows_sampled += 1
                sampled_sheet_rows += 1
                columns_sampled = max(columns_sampled, min(len(cells), MAX_COLUMNS))
                rendered = []
                for cell in cells:
                    value = cell.value
                    if value is None:
                        continue
                    cells_sampled += 1
                    if cell.data_type == "f":
                        formulas_sampled += 1
                    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                        has_long_cells = True
                    rendered.append(f"{cell.coordinate}={_json_cell(value)}")
                if rendered and not builder.add(" | ".join(rendered)):
                    stopped = True
                    break
            if builder.truncated:
                break
            if sampled_sheet_rows >= row_limit and reported_rows > sampled_sheet_rows:
                stopped = True
    except Exception as exc:
        raise ValueError("Não foi possível ler a planilha XLSX; arquivos corrompidos ou protegidos não são aceitos.") from exc
    finally:
        workbook.close()

    if stopped:
        warnings.append(f"A amostra foi limitada a até {MAX_TOTAL_ROWS} linhas e {MAX_COLUMNS} colunas por aba.")
    if builder.truncated:
        warnings.append(f"O contexto foi limitado a {MAX_CONTEXT_CHARS:,} caracteres.")
    if has_long_cells:
        warnings.append(f"Células maiores que {MAX_CELL_CHARS} caracteres foram abreviadas.")
    context = builder.text()
    return {
        "context": context,
        "preview": context[:MAX_PREVIEW_CHARS],
        "sheets": selected_names,
        "statistics": {
            "size_bytes": len(content), "sheets": len(sheet_names), "rows_sampled": rows_sampled,
            "columns_sampled": columns_sampled, "cells_sampled": cells_sampled,
            "formulas_sampled": formulas_sampled,
            "truncated": stopped or builder.truncated or has_long_cells or len(sheet_names) > MAX_SHEETS,
        },
        "warnings": warnings,
    }


def _root(runs_root=None):
    return Path(runs_root if runs_root is not None else RUNS) / "attachments"


def save_attachment(name, content_base64, runs_root=None):
    """Validate, extract, and persist one attachment under a random server-side ID."""
    name, suffix = _validate_name(name)
    content = _decode_base64(content_base64)
    kind = suffix[1:]
    extracted = _extract_xlsx(name, content) if suffix == ".xlsx" else _extract_csv(name, suffix, content)
    attachment_id = str(uuid.uuid4())
    directory = _root(runs_root) / attachment_id
    directory.mkdir(parents=True, exist_ok=False)
    original = directory / ("original" + suffix)
    try:
        original.write_bytes(content)
        metadata = {
            "id": attachment_id, "name": name, "kind": kind,
            "sha256": hashlib.sha256(content).hexdigest(),
            "context": extracted["context"], "preview": extracted["preview"],
            "sheets": extracted["sheets"], "statistics": extracted["statistics"],
            "warnings": extracted["warnings"],
        }
        write_json(directory / "manifest.json", metadata)
    except Exception:
        original.unlink(missing_ok=True)
        (directory / "manifest.json").unlink(missing_ok=True)
        directory.rmdir()
        raise
    return metadata


def load_attachments(attachment_ids, runs_root=None):
    """Resolve IDs to stored metadata without accepting client-supplied paths or context."""
    if not isinstance(attachment_ids, list) or len(attachment_ids) > MAX_ATTACHMENTS_PER_MESSAGE:
        raise ValueError(f"Anexe no máximo {MAX_ATTACHMENTS_PER_MESSAGE} arquivos por mensagem.")
    if any(not isinstance(attachment_id, str) for attachment_id in attachment_ids):
        raise ValueError("Identificador de anexo inválido.")
    if len(set(attachment_ids)) != len(attachment_ids):
        raise ValueError("Um arquivo foi selecionado mais de uma vez.")
    root = _root(runs_root)
    attachments = []
    for attachment_id in attachment_ids:
        if not isinstance(attachment_id, str) or not ATTACHMENT_ID.fullmatch(attachment_id):
            raise ValueError("Identificador de anexo inválido.")
        directory = root / attachment_id
        if not directory.is_dir() or directory.resolve().parent != root.resolve():
            raise ValueError("Anexo não encontrado. Selecione o arquivo novamente.")
        try:
            record = read_json(directory / "manifest.json")
        except (OSError, ValueError):
            record = None
        if not isinstance(record, dict) or record.get("id") != attachment_id:
            raise ValueError("Anexo não encontrado. Selecione o arquivo novamente.")
        name = record.get("name")
        kind = record.get("kind")
        context = record.get("context")
        if (not isinstance(name, str) or kind not in {"xlsx", "csv", "tsv"}
                or not isinstance(context, str) or not context or len(context) > MAX_CONTEXT_CHARS):
            raise ValueError("Os dados do anexo estão inválidos. Selecione o arquivo novamente.")
        attachments.append({
            "id": attachment_id, "name": name, "kind": kind, "context": context,
            "preview": str(record.get("preview", ""))[:MAX_PREVIEW_CHARS],
            "sheets": record.get("sheets", []) if isinstance(record.get("sheets", []), list) else [],
            "statistics": record.get("statistics", {}) if isinstance(record.get("statistics", {}), dict) else {},
            "warnings": record.get("warnings", []) if isinstance(record.get("warnings", []), list) else [],
        })
    return attachments
